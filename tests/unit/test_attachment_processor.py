"""Unit tests for enp/attachment_processor.py: reading real (tiny) files of every supported type.

The files are built in memory, so nothing is downloaded and no Gmail account is needed.
"""
import base64
import io
import zipfile

import pytest

from enp import attachment_processor as ap


# ── Builders for small real files ─────────────────────────────────────────────

def make_pdf(text="Hello PDF world"):
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 18 Tf 20 80 Td ({text}) Tj ET"
    objects[3] = f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def make_docx():
    import docx

    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "ACME Letterhead"
    document.add_paragraph("Offer letter for the backend role")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Salary"
    table.rows[0].cells[1].text = "12 LPA"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_zip(members):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return buffer.getvalue()


def make_pptx():
    slide = ('<p:sld><a:p><a:r><a:t>Quarterly &amp; plan</a:t></a:r></a:p>'
             '<a:p><a:r><a:t>Line one</a:t></a:r><a:br/><a:r><a:t>Line two</a:t></a:r></a:p></p:sld>')
    notes = '<p:notes><a:p><a:r><a:t>Speaker note</a:t></a:r></a:p></p:notes>'
    return make_zip({"ppt/slides/slide2.xml": slide.replace("Quarterly", "Second"), "ppt/slides/slide10.xml": slide.replace("Quarterly", "Tenth"),
                     "ppt/slides/slide1.xml": slide, "ppt/notesSlides/notesSlide1.xml": notes})


def make_xlsx():
    shared = "<sst><si><t>Name</t></si><si><t>Priya</t></si><si><r><t>Dev</t></r><rPh><t>ignored</t></rPh></si></sst>"
    sheet = ('<worksheet><sheetData><row><c t="s"><v>0</v></c><c t="s"><v>1</v></c></row>'
             '<row><c t="inlineStr"><is><t>Score</t></is></c><c><v>42</v></c><c t="s"><v>99</v></c><c/></row></sheetData></worksheet>')
    workbook = '<workbook><sheets><sheet name="Results &amp; more" r:id="rId1"/></sheets></workbook>'
    rels = '<Relationships><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>'
    return make_zip({"xl/sharedStrings.xml": shared, "xl/worksheets/sheet1.xml": sheet,
                     "xl/workbook.xml": workbook, "xl/_rels/workbook.xml.rels": rels})


# ── Extractors ────────────────────────────────────────────────────────────────

def test_pdf_text_is_extracted():
    assert "Hello PDF world" in ap.extract_text(make_pdf(), "application/pdf", "a.pdf")


def test_pdf_by_extension_when_gmail_says_octet_stream():
    """Shows: Gmail often labels files application/octet-stream, so the extension decides."""
    assert "Hello PDF" in ap.extract_text(make_pdf(), "application/octet-stream", "report.PDF")


def test_a_pdf_with_no_text_layer_goes_to_ocr(monkeypatch):
    monkeypatch.setattr(ap, "_extract_pdf", lambda data, max_chars: None)
    monkeypatch.setattr(ap, "_ocr_pdf", lambda data, max_chars: "scanned words")
    assert ap.extract_text(b"%PDF", "application/pdf", "scan.pdf") == "scanned words"


def test_a_broken_pdf_gives_none_instead_of_crashing():
    assert ap._extract_pdf(b"this is not a pdf", 1000) is None
    assert ap._extract_pdf_plumber(b"nope", 1000) is None


def test_pdf_reading_stops_at_the_character_cap(monkeypatch):
    class Page:
        def extract_text(self):
            return "x" * 100

    class Reader:
        is_encrypted = False
        pages = [Page() for _ in range(50)]

        def __init__(self, stream):
            pass

    monkeypatch.setattr("pypdf.PdfReader", Reader)
    assert len(ap._extract_pdf(b"", max_chars=250)) == 300 + 2   # 3 pages joined by newlines


def test_an_encrypted_pdf_that_needs_a_password_is_skipped(monkeypatch):
    class Reader:
        is_encrypted = True
        pages = []

        def __init__(self, stream):
            pass

        def decrypt(self, password):
            return 0

    monkeypatch.setattr("pypdf.PdfReader", Reader)
    assert ap._extract_pdf(b"", 1000) is None


def test_docx_text_includes_header_paragraphs_and_tables():
    text = ap.extract_text(make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "offer.docx")
    assert "ACME Letterhead" in text and "Offer letter" in text and "Salary | 12 LPA" in text
    assert ap._extract_docx(b"not a docx") is None


def test_plain_text_decoding_handles_boms_and_latin1():
    assert ap._decode_text("héllo".encode("utf-8-sig")) == "héllo"
    assert ap._decode_text("héllo".encode("utf-16")) == "héllo"
    assert ap._decode_text("héllo".encode("latin-1")) == "héllo"
    assert ap.extract_text(b"line1\r\nline2  ", "text/plain", "n.txt") == "line1\nline2"
    assert ap.extract_text(b"   ", "text/plain", "n.txt") is None
    assert ap.extract_text(b"a,b\n1,2", "application/octet-stream", "data.csv") == "a,b\n1,2"


def test_html_attachment_is_stripped_to_text():
    assert ap.extract_text(b"<html><body><p>Hi</p><script>x()</script></body></html>", "text/html", "p.html") == "Hi"


def test_pptx_slides_are_read_in_numeric_order_with_notes_last():
    text = ap.extract_text(make_pptx(), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "deck.pptx")
    assert text.index("Quarterly") < text.index("Second") < text.index("Tenth") < text.index("Speaker note")
    assert "Quarterly & plan" in text and "Line one\nLine two" in text


def test_xlsx_rows_become_lines_with_the_sheet_name():
    text = ap.extract_text(make_xlsx(), "application/vnd.ms-excel.sheet.macroEnabled.12", "book.xlsm")
    assert text.startswith("[Results & more]")
    assert "Name | Priya" in text and "Score | 42" in text
    assert "ignored" not in text            # phonetic hints are dropped


def test_office_files_that_are_not_zips_or_too_large_give_none(monkeypatch):
    assert ap._extract_pptx(b"not a zip", 1000) is None
    assert ap._extract_xlsx(b"not a zip", 1000) is None
    monkeypatch.setattr(ap, "OFFICE_MEMBER_MAX_BYTES", 10)
    assert ap._extract_pptx(make_pptx(), 1000) is None            # every member is over the cap


def test_eml_attachment_keeps_headers_and_body():
    from email.message import EmailMessage

    message = EmailMessage()
    message["Subject"], message["From"], message["To"] = "Offer", "hr@acme.com", "me@x.com"
    message.set_content("Please confirm by Friday.")
    text = ap.extract_text(message.as_bytes(), "message/rfc822", "fwd.eml")
    assert text.startswith("Subject: Offer") and "Please confirm by Friday." in text
    html_only = EmailMessage()
    html_only.set_content("<p>Hello <b>there</b></p>", subtype="html")
    assert "Hello" in ap.extract_text(html_only.as_bytes(), "message/rfc822", "x.eml")
    assert ap._extract_eml(None) is None


def test_unsupported_types_return_none():
    """Shows (decision A22): legacy .doc, RTF and zip files are not read."""
    for mime, name in [("application/msword", "old.doc"), ("text/rtf", "x.rtf"), ("application/zip", "a.zip"), ("", "")]:
        assert ap.extract_text(b"data", mime, name) is None


def test_images_are_only_read_when_ocr_is_ready(monkeypatch):
    monkeypatch.setattr(ap, "_ocr_ready", lambda: False)
    assert ap.extract_text(b"\x89PNG", "image/png", "scan.png") is None
    assert ap.can_extract_text("image/png", "scan.png") is False
    monkeypatch.setattr(ap, "_ocr_ready", lambda: True)
    assert ap.can_extract_text("image/png", "scan.png") is True
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (10, 10), "white").save(buffer, format="PNG")
    monkeypatch.setattr(ap, "_ocr_image", lambda image: "text from picture")
    assert ap.extract_text(buffer.getvalue(), "image/png", "scan.png") == "text from picture"
    assert ap._extract_image_ocr(b"not an image") is None


def test_ocr_is_off_without_the_tesseract_program(monkeypatch):
    ap._ocr_ready.cache_clear()
    monkeypatch.delenv("TESSERACT_CMD", raising=False)
    monkeypatch.setattr(ap.shutil, "which", lambda name: None)
    monkeypatch.setattr(ap.os.path, "isfile", lambda path: False)
    import pytesseract

    monkeypatch.setattr(pytesseract, "get_tesseract_version", lambda: (_ for _ in ()).throw(OSError("not installed")))
    try:
        assert ap._ocr_ready() is False
    finally:
        ap._ocr_ready.cache_clear()


def test_scanned_pdf_ocr_stops_after_the_page_limit(monkeypatch):
    class Image:
        def to_pil(self):
            return "image"

    class Page:
        def render(self, scale):
            return Image()

    class Document:
        closed = False

        def __init__(self, data):
            self.pages = [Page() for _ in range(25)]

        def __len__(self):
            return len(self.pages)

        def __getitem__(self, index):
            return self.pages[index]

        def close(self):
            Document.closed = True

    import pypdfium2

    monkeypatch.setattr(pypdfium2, "PdfDocument", Document)
    monkeypatch.setattr(ap, "_ocr_ready", lambda: True)
    monkeypatch.setattr(ap, "_ocr_image", lambda image: "page text")
    text = ap._ocr_pdf(b"%PDF", max_chars=10_000)
    assert text.count("page text") == ap.OCR_PDF_MAX_PAGES and Document.closed
    monkeypatch.setattr(ap, "_ocr_ready", lambda: False)
    assert ap._ocr_pdf(b"%PDF", 1000) is None


def test_can_extract_text_mirrors_the_routing():
    for mime, name in [("application/pdf", "a.pdf"), ("", "x.docx"), ("", "s.pptx"), ("", "s.xlsx"), ("", "m.eml"),
                       ("text/html", "p.html"), ("text/plain", "n.txt"), ("application/octet-stream", "d.csv")]:
        assert ap.can_extract_text(mime, name), name
    for mime, name in [("application/msword", "old.doc"), ("text/rtf", "x.rtf"), ("application/zip", "z.zip")]:
        assert not ap.can_extract_text(mime, name), name


# ── Gmail part traversal ──────────────────────────────────────────────────────

def b64(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).decode().rstrip("=")


def test_base64url_decoding_accepts_missing_padding():
    assert ap._b64url_decode(b64(b"hello!?")) == b"hello!?"


def test_walk_parts_flattens_nested_parts_in_message_order():
    payload = {"filename": "", "parts": [{"filename": "a.txt", "parts": [{"filename": "inner.txt"}]}, {"filename": "b.txt"}]}
    assert [p["filename"] for p in ap._walk_parts(payload)] == ["", "a.txt", "inner.txt", "b.txt"]


def test_same_named_attachments_are_told_apart_by_size():
    """Shows (A22.3): two files called report.pdf are separated by their byte size."""
    parts = [{"filename": "report.pdf", "body": {"size": 10}}, {"filename": "report.pdf", "body": {"size": 99}}]
    assert ap._select_attachment_parts(parts, "report.pdf", 99) == [parts[1]]
    assert ap._select_attachment_parts(parts, "report.pdf", 0) == parts
    assert ap._select_attachment_parts(parts, "report.pdf", 5) == parts      # no size match: try them all
    assert ap._select_attachment_parts(parts, "other.pdf", 10) == []


class FakeGmailService:
    def __init__(self, message, blobs=None):
        self.message, self.blobs, self.calls = message, blobs or {}, 0

    def users(self):
        return self

    def messages(self):
        return self

    def attachments(self):
        return self

    def get(self, **kwargs):
        self.calls += 1
        self.last = kwargs
        return self

    def execute(self):
        if "messageId" in self.last:
            return self.blobs.get(self.last["id"], {})
        return self.message


def test_download_inline_and_separately_fetched_attachments():
    inline = {"payload": {"parts": [{"filename": "a.txt", "body": {"data": b64(b"inline text"), "size": 11}}]}}
    assert ap._download_attachment_data(FakeGmailService(inline), "e1", "a.txt") == b"inline text"
    remote = {"payload": {"parts": [{"filename": "b.txt", "body": {"attachmentId": "att1", "size": 5}}]}}
    service = FakeGmailService(remote, {"att1": {"data": b64(b"fetched")}})
    assert ap._download_attachment_data(service, "e1", "b.txt", 5) == b"fetched"
    assert ap._download_attachment_data(FakeGmailService(remote), "e1", "b.txt") is None
    assert ap._download_attachment_data(FakeGmailService({"payload": {}}), "e1", "missing.txt") is None


# ── Ingestion-time extraction (A4) ────────────────────────────────────────────

def service_with(filename, content):
    message = {"payload": {"parts": [{"filename": filename, "body": {"data": b64(content), "size": len(content)}}]}}
    return FakeGmailService(message)


def test_extract_full_text_reads_supported_files_only():
    service = service_with("n.txt", b"offer details")
    assert ap.extract_full_text(service, "e1", {"filename": "n.txt", "mime_type": "text/plain", "file_size": 13}) == "offer details"
    assert ap.extract_full_text(service, "e1", {"filename": "a.zip", "mime_type": "application/zip"}) is None
    assert service.calls == 1               # the zip was never downloaded


def test_extract_full_text_skips_files_over_ten_megabytes_without_downloading():
    service = service_with("big.txt", b"x")
    big = {"filename": "big.txt", "mime_type": "text/plain", "file_size": ap.MAX_INGEST_ATTACHMENT_BYTES + 1}
    assert ap.extract_full_text(service, "e1", big) is None and service.calls == 0


def test_extract_full_text_cuts_at_the_stored_text_limit(monkeypatch):
    monkeypatch.setattr(ap, "MAX_EXTRACTED_TEXT_CHARS", 10)
    service = service_with("n.txt", b"x" * 100)
    assert ap.extract_full_text(service, "e1", {"filename": "n.txt", "mime_type": "text/plain"}) == "x" * 10


# ── Tier-2 selection and orchestration ────────────────────────────────────────

def test_attachments_qualify_by_filename_match_or_small_size():
    assert ap._qualifies({"filename": "Salary_Slip.pdf", "file_size": 10_000_000}, "salary slip october")
    assert not ap._qualifies({"filename": "big.pdf", "file_size": 10_000_000}, "the pdf file")   # filler words never match
    assert ap._qualifies({"filename": "small.pdf", "file_size": 1000}, "")
    assert ap._qualifies({"filename": "unknown.pdf", "file_size": 0}, "")
    assert not ap._qualifies({"filename": "big.pdf", "file_size": 600 * 1024}, "")


@pytest.fixture
def tier2(monkeypatch, fake_redis):
    from ste import redis_manager

    monkeypatch.setattr(redis_manager, "get_sync_client", lambda: fake_redis)
    state = {"downloads": 0, "marked": []}
    rows = [{"attachment_id": "a1", "filename": "offer.txt", "mime_type": "text/plain", "file_size": 100,
             "lightweight_extract": "offer.txt | text/plain | 100"},
            {"attachment_id": "a2", "filename": "scan.png", "mime_type": "image/png", "file_size": 100,
             "lightweight_extract": "scan.png | image/png | 100"}]
    monkeypatch.setattr(ap.postgresql_manager, "fetchone", lambda q, p=None: {"email_id": "e1"})
    monkeypatch.setattr(ap.postgresql_manager, "fetchall", lambda q, p=None: rows)
    monkeypatch.setattr(ap.postgresql_manager, "execute", lambda q, p=None: state["marked"].append(p["attachment_id"]))

    def download(service, email_id, filename, size=0):
        state["downloads"] += 1
        return b"Salary is 12 LPA" if filename == "offer.txt" else None

    monkeypatch.setattr(ap, "_download_attachment_data", download)
    monkeypatch.setattr("ingestion.gmail.gmail_api.authenticate_gmail", lambda: "service")
    return state


def test_tier2_extracts_caches_and_falls_back_to_metadata(tier2):
    text = ap.fetch_attachment_text_for_memory("m1", "offer")
    assert "── offer.txt ──\nSalary is 12 LPA" in text
    assert "(metadata only) scan.png | image/png | 100" in text      # unreadable -> the answer degrades, never disappears
    assert tier2["marked"] == ["a1"]
    first_downloads = tier2["downloads"]
    again = ap.fetch_attachment_text_for_memory("m1", "offer")        # second call is served from Redis
    assert "Salary is 12 LPA" in again and tier2["downloads"] == first_downloads + 1   # only the unreadable one is retried


def test_tier2_respects_the_total_character_cap(tier2, monkeypatch):
    monkeypatch.setattr(ap, "MAX_CHARS_TOTAL", 20)
    text = ap.fetch_attachment_text_for_memory("m1", "offer")
    assert "Salary is 12 LPA" in text
    assert text.endswith("── scan.png ──\n(met")             # the second block is cut to the 4 characters left


def test_tier2_never_raises(monkeypatch):
    monkeypatch.setattr(ap.postgresql_manager, "fetchone", lambda q, p=None: (_ for _ in ()).throw(RuntimeError("db down")))
    assert ap.fetch_attachment_text_for_memory("m1") is None
    monkeypatch.setattr(ap.postgresql_manager, "fetchone", lambda q, p=None: None)
    assert ap.fetch_attachment_text_for_memory("m1") is None
    monkeypatch.setattr(ap.postgresql_manager, "fetchone", lambda q, p=None: {"email_id": "e"})
    monkeypatch.setattr(ap.postgresql_manager, "fetchall", lambda q, p=None: [])
    assert ap.fetch_attachment_text_for_memory("m1") is None


def test_tier2_survives_gmail_and_redis_failures(tier2, monkeypatch):
    monkeypatch.setattr("ingestion.gmail.gmail_api.authenticate_gmail", lambda: (_ for _ in ()).throw(RuntimeError("no token")))
    from ste import redis_manager

    monkeypatch.setattr(redis_manager, "get_sync_client", lambda: (_ for _ in ()).throw(ConnectionError("down")))
    text = ap.fetch_attachment_text_for_memory("m1", "offer")
    assert "(metadata only) offer.txt" in text           # no Gmail, no Redis: still answers from Tier 1


def test_marking_failure_is_swallowed(monkeypatch):
    monkeypatch.setattr(ap.postgresql_manager, "execute", lambda q, p=None: (_ for _ in ()).throw(RuntimeError("down")))
    ap._mark_extracted("a1")        # visibility flags only; must not raise
